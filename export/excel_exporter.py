from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter
from models import MonthPlan, ShiftType, is_effective
import calendar


def _hex_to_rgb(hex_color: str) -> tuple:
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))


def export_excel(plan: MonthPlan, folder: str | Path) -> str:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)

    year, month = plan.year, plan.month
    days_in_month = calendar.monthrange(year, month)[1]

    excel_file = folder / f"plan_{year}_{month:02d}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Dienstplan"

    # Header-Zeile: Tage + Wochentag
    ws.cell(1, 1, "Assistent")
    for day in range(1, days_in_month + 1):
        import datetime as dt
        d = dt.date(year, month, day)
        weekday = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"][d.weekday()]
        ws.cell(1, day + 1, f"{day}\n{weekday}")

    # Daten
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row, assistant in enumerate(plan.assistants, 2):
        ws.cell(row, 1, assistant.name)

        for day in range(1, days_in_month + 1):
            col = day + 1
            entries = plan.schedule.get(day, [])
            matching = [e for e in entries
                        if e.assistant_id == assistant.id and is_effective(e)]

            cell = ws.cell(row, col)
            if matching:
                entry = matching[0]
                if entry.shift_type == ShiftType.FULL:
                    cell.value = "VOLL"
                elif entry.shift_type == ShiftType.HALF_MORNING:
                    cell.value = "VM"
                elif entry.shift_type == ShiftType.HALF_AFTERNOON:
                    cell.value = "NM"
                elif entry.shift_type == ShiftType.ON_CALL:
                    cell.value = "RB"

                # Farbe
                try:
                    rgb = _hex_to_rgb(assistant.color)
                    fill = PatternFill(start_color=f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}", end_color=f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}", fill_type="solid")
                    cell.fill = fill
                except:
                    pass

            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

            # Wochenende Hintergrund
            import datetime as dt
            d = dt.date(year, month, day)
            if d.weekday() >= 5 and not matching:
                cell.fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

    # Zusammenfassung-Zeile
    summary_row = len(plan.assistants) + 3
    ws.cell(summary_row, 1, "ZUSAMMENFASSUNG")

    summary_ws = wb.create_sheet("Zusammenfassung")
    summary_ws.cell(1, 1, "Assistent")
    summary_ws.cell(1, 2, "Dienste gesamt")
    summary_ws.cell(1, 3, "VOLL")
    summary_ws.cell(1, 4, "VM")
    summary_ws.cell(1, 5, "NM")
    summary_ws.cell(1, 6, "RB")

    for row, assistant in enumerate(plan.assistants, 2):
        summary_ws.cell(row, 1, assistant.name)

        full_count = 0
        vm_count = 0
        nm_count = 0
        rb_count = 0

        for entries in plan.schedule.values():
            for e in entries:
                if e.assistant_id == assistant.id and is_effective(e):
                    if e.shift_type == ShiftType.FULL:
                        full_count += 1
                    elif e.shift_type == ShiftType.HALF_MORNING:
                        vm_count += 1
                    elif e.shift_type == ShiftType.HALF_AFTERNOON:
                        nm_count += 1
                    elif e.shift_type == ShiftType.ON_CALL:
                        rb_count += 1

        # Gewichtete Dienstzahl (VOLL=1, VM/NM=0,5); RB zaehlt separat
        summary_ws.cell(row, 2, full_count + 0.5 * (vm_count + nm_count))
        summary_ws.cell(row, 3, full_count)
        summary_ws.cell(row, 4, vm_count)
        summary_ws.cell(row, 5, nm_count)
        summary_ws.cell(row, 6, rb_count)

    # Breiten anpassen
    ws.column_dimensions["A"].width = 20
    for col in range(2, days_in_month + 2):
        ws.column_dimensions[get_column_letter(col)].width = 12

    summary_ws.column_dimensions["A"].width = 20
    for col in range(2, 7):
        summary_ws.column_dimensions[get_column_letter(col)].width = 15

    wb.save(excel_file)
    return str(excel_file)
